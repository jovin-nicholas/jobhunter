import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import docx

from jobhunter import resumes
from jobhunter.errors import SettingsError
from jobhunter.models import Job
from jobhunter.resumes import header_name, anonymize, extract_text, load_resumes
from jobhunter.settings import ResumeSettings
from tests.helpers import RESUME_TEXT, make_pdf

FILLER = " Built and shipped backend services, APIs, dashboards and data pipelines for production users." * 3


def job(title="Software Engineer", description=""):
    return Job("t_1", title, "Acme", "Austin, TX", "https://example.com/1", description=description)


class ResumeTestCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.folder = self.tmp / "resumes"
        self.folder.mkdir()
        self.cache = self.tmp / "cache"

    def write(self, name, text=RESUME_TEXT):
        (self.folder / name).write_text(text, encoding="utf-8")

    def settings(self, default="backend.txt", versions=None, names=None):
        return ResumeSettings(self.folder, default, versions or {}, names or [])


class TestExtractText(ResumeTestCase):
    def test_txt_md_pdf_and_docx(self):
        self.write("a.txt", "plain text")
        self.write("b.md", "# markdown")
        make_pdf(self.folder / "c.pdf", ["John Doe - Backend Engineer", "Built REST APIs in Java."])
        document = docx.Document()
        document.add_paragraph("John Doe - Backend Engineer")
        document.add_paragraph("Python and React")
        document.save(self.folder / "d.docx")
        self.assertEqual(extract_text(self.folder / "a.txt"), "plain text")
        self.assertEqual(extract_text(self.folder / "b.md"), "# markdown")
        self.assertIn("Built REST APIs in Java.", extract_text(self.folder / "c.pdf"))
        self.assertEqual(extract_text(self.folder / "d.docx"), "John Doe - Backend Engineer\nPython and React")


class TestNames(ResumeTestCase):
    def test_header_name_removed_everywhere_including_email_and_url(self):
        self.write("backend.txt")
        text = load_resumes(self.settings(), self.cache, env={}).resumes["backend.txt"].text
        self.assertTrue(text.startswith("the candidate — New Grad"), text[:40])
        self.assertNotIn("john", text.lower())

    def test_a_name_alone_on_the_first_line(self):
        for first in ("John Doe", "Mary-Jane O'Neil", "J. R. Doe", "JOHN DOE"):
            self.assertEqual(header_name(f"{first}\njohn.doe@example.com | 555-0100\nSkills: Python"), first, first)
        for first in ("Software Engineer", "Resume", "Curriculum Vitae", "john.doe@example.com",
                      "John Doe 555-0100", "Senior Backend Software Engineer Resume Draft"):
            self.assertEqual(header_name(f"{first}\nSkills: Python"), "", first)

    def test_the_resume_set_says_which_names_it_removes(self):
        self.write("backend.txt", "Skills: Python and Go" + FILLER)
        self.assertEqual(load_resumes(self.settings(), self.cache, env={}).names, [])
        self.write("backend.txt", "John Doe\nSkills: Python and Go" + FILLER)
        loaded = load_resumes(self.settings(), self.cache, env={})
        self.assertEqual(loaded.names, ["John Doe"])
        self.assertNotIn("John", loaded.resumes["backend.txt"].text)

    def test_accented_names(self):
        text = anonymize("José García — SDE\nContact: jose@example.com\nJOSÉ GARCÍA led the team.", ["José García"])
        self.assertNotIn("José", text)
        self.assertNotIn("JOSÉ", text)
        self.assertIn("the candidate led the team", text)

    def test_first_name_alone_is_not_replaced(self):
        self.assertEqual(anonymize("Will Smith — Engineer\nI will build things.", ["Will Smith"]),
                         "the candidate — Engineer\nI will build things.")

    def test_settings_names_beat_env_which_beats_header(self):
        self.write("backend.txt", "Jane Roe — SDE\nJohnny Doe wrote this. John Doe reviewed it." + FILLER)

        def load(settings, env):
            return load_resumes(settings, self.cache, env=env).resumes["backend.txt"].text

        from_settings = load(self.settings(names=["Johnny Doe"]), {"CANDIDATE_NAMES": "John Doe"})
        self.assertIn("the candidate wrote this. John Doe reviewed", from_settings)
        from_env = load(self.settings(), {"CANDIDATE_NAMES": "John Doe"})
        self.assertIn("Jane Roe — SDE\nJohnny Doe wrote this. the candidate reviewed", from_env)
        from_header = load(self.settings(), {})
        self.assertTrue(from_header.startswith("the candidate — SDE\nJohnny Doe wrote this. John Doe reviewed"))


class TestLoading(ResumeTestCase):
    def test_short_text_is_a_problem(self):
        self.write("backend.txt", "John Doe — SDE\nToo short.")
        with self.assertRaises(SettingsError) as caught:
            load_resumes(self.settings(), self.cache, env={})
        self.assertIn("characters of text", caught.exception.problems[0])

    def test_unsupported_format_is_a_problem(self):
        self.write("resume.pages", RESUME_TEXT)
        with self.assertRaises(SettingsError) as caught:
            load_resumes(self.settings(default="resume.pages"), self.cache, env={})
        self.assertIn("unsupported format", caught.exception.problems[0])

    def test_unrelated_files_in_the_folder_are_ignored(self):
        self.write("backend.txt")
        (self.folder / ".DS_Store").write_bytes(b"\x00\x01binary")
        self.write("notes.pages", "not a resume")
        self.assertEqual(list(load_resumes(self.settings(), self.cache, env={}).resumes), ["backend.txt"])

    def test_pdf_text_is_cached_until_the_file_changes(self):
        make_pdf(self.folder / "backend.pdf", ["John Doe - Backend Engineer", RESUME_TEXT.replace("\n", " ").replace("—", "-")])
        settings = self.settings(default="backend.pdf")
        with patch.object(resumes, "extract_text", wraps=resumes.extract_text) as spy:
            load_resumes(settings, self.cache, env={})
            load_resumes(settings, self.cache, env={})
        self.assertEqual(spy.call_count, 1)


class TestPick(ResumeTestCase):
    def setUp(self):
        super().setUp()
        for name in ("backend.txt", "fullstack.txt", "data.txt"):
            self.write(name)
        self.set = load_resumes(self.settings(versions={"fullstack.txt": ["react", "node.js"],
                                                        "data.txt": ["spark", "react"]}), self.cache, env={})

    def test_most_keyword_matches_wins(self):
        self.assertEqual(self.set.pick(job("Full Stack Engineer", "React and Node.js")).id, "fullstack.txt")

    def test_tie_and_no_match_use_the_default(self):
        self.assertEqual(self.set.pick(job(description="React dashboards")).id, "backend.txt")
        self.assertEqual(self.set.pick(job(description="Java services")).id, "backend.txt")


if __name__ == "__main__":
    unittest.main()
